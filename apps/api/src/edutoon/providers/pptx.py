"""PowerPoint (``.pptx``) text extraction - the only module allowed to import ``python-pptx``
(rule 3).

A deck is read into :class:`~edutoon.domain.deck.DeckSlide` structures, not flat text: title,
body paragraphs, bullets with their level, numbered steps, tables and speaker notes are each kept
as what they are (``domain.deck``), because the lesson generator narrates them differently.

What is read: text in any shape (grouped shapes included), tables, and speaker notes. What is not:
pictures, charts and SmartArt carry no text this library exposes, so they contribute nothing -
a slide that is only a picture has no content and is left out of the lesson. Slides the author
hid (``show="0"``) are skipped, as they would be in a presentation; every slide keeps the number
it has in the deck, so a citation points at the slide the author sees.

A ``.pptx`` is a zip file, so it is checked before it is opened: the size it would expand to and
the number of parts are bounded, the same defence the PDF reader has (``providers.pdf``).
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from edutoon.domain.deck import DeckBlock, DeckSlide, clean_text, normalise_slide

# A real deck with large images expands to a few hundred megabytes at most.
MAX_UNCOMPRESSED_BYTES = 600 * 1024 * 1024
MAX_ZIP_ENTRIES = 20_000

_NOT_CONTENT = {
    PP_PLACEHOLDER.SLIDE_NUMBER,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.HEADER,
}
_TITLES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_SUBTITLES = {PP_PLACEHOLDER.SUBTITLE}

_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class PptxParseError(Exception):
    """``content`` cannot be read as a PowerPoint presentation at all (corrupted, encrypted,
    not a ``.pptx``, or a zip that would expand unreasonably)."""


class PptxTooManySlidesError(PptxParseError):
    """A well-formed deck with more slides than the caller allows. Raised before any slide's
    text is read."""

    def __init__(self, slide_count: int, max_slides: int) -> None:
        super().__init__(f"The presentation has {slide_count} slides; the limit is {max_slides}.")
        self.slide_count = slide_count
        self.max_slides = max_slides


@dataclass(frozen=True, slots=True)
class ParsedDeck:
    title: str | None
    slides: tuple[DeckSlide, ...]
    slide_count: int
    """How many slides the deck has, hidden ones included."""


def looks_like_pptx(content: bytes) -> bool:
    """A cheap check before parsing: a zip archive holding a presentation part."""
    if not content.startswith(b"PK\x03\x04"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return "ppt/presentation.xml" in archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return False


def _check_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise PptxParseError("The presentation has an unreasonable number of parts.")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise PptxParseError("The presentation expands to an unreasonable size.")
    except zipfile.BadZipFile as exc:
        raise PptxParseError("The file is not a valid .pptx presentation.") from exc


def _placeholder_type(shape: BaseShape) -> PP_PLACEHOLDER | None:
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except (ValueError, KeyError):  # a placeholder whose type the library cannot resolve
        return None


def _paragraph_kind(paragraph_xml: object, *, in_body: bool) -> str:
    """``bullet``, ``numbered`` or ``paragraph`` for one paragraph.

    An explicit bullet character or auto-number says so; ``buNone`` says it is plain text; with
    neither, a paragraph in a body placeholder is a bullet (the template's default) and one in a
    free text box is a paragraph.
    """
    properties = paragraph_xml.find(f"{_A_NS}pPr")  # type: ignore[attr-defined]
    if properties is not None:
        if properties.find(f"{_A_NS}buNone") is not None:
            return "paragraph"
        if properties.find(f"{_A_NS}buAutoNum") is not None:
            return "numbered"
        if (
            properties.find(f"{_A_NS}buChar") is not None
            or properties.find(f"{_A_NS}buBlip") is not None
        ):
            return "bullet"
    return "bullet" if in_body else "paragraph"


def _text_blocks(shape: BaseShape, *, in_body: bool) -> list[DeckBlock]:
    blocks: list[DeckBlock] = []
    frame = getattr(shape, "text_frame", None) if shape.has_text_frame else None
    if frame is None:
        return blocks
    for paragraph in frame.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        kind = _paragraph_kind(paragraph._p, in_body=in_body)
        blocks.append(
            DeckBlock(
                kind=kind,  # type: ignore[arg-type]
                text=text,
                level=paragraph.level if kind in {"bullet", "numbered"} else 0,
            )
        )
    return blocks


def _walk(shapes: object) -> list[BaseShape]:
    """Every shape, group members included, in document order."""
    found: list[BaseShape] = []
    for shape in shapes:  # type: ignore[attr-defined]
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found.extend(_walk(shape.shapes))
        else:
            found.append(shape)
    return found


def _reading_order(shapes: list[BaseShape]) -> list[BaseShape]:
    """Top to bottom, then left to right - when every shape says where it is; otherwise the
    order they are stored in, which is the order the author added them."""
    try:
        positions = [(shape.top, shape.left) for shape in shapes]
    except (AttributeError, ValueError):
        return shapes
    if any(top is None or left is None for top, left in positions):
        return shapes
    # Shapes whose tops are within a line of each other count as one row, ordered left to right.
    line = 228_600  # a quarter of an inch, in EMU
    return sorted(shapes, key=lambda shape: (round((shape.top or 0) / line), shape.left or 0))


def _slide_title(slide: Slide) -> tuple[str | None, BaseShape | None]:
    shape = slide.shapes.title
    if shape is not None and shape.has_text_frame:
        title = clean_text(shape.text_frame.text)
        if title:
            return title, shape
    return None, shape


def _slide_notes(slide: Slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return clean_text(frame.text) if frame is not None else ""


def _read_slide(slide: Slide, *, index: int, deck_title: str | None) -> DeckSlide:
    title, title_shape = _slide_title(slide)
    blocks: list[DeckBlock] = []
    for shape in _reading_order(_walk(slide.shapes)):
        if title_shape is not None and shape.shape_id == title_shape.shape_id:
            continue
        placeholder = _placeholder_type(shape)
        if placeholder in _NOT_CONTENT:
            continue
        if placeholder in _TITLES:
            # A second title-like placeholder (a vertical or secondary title) reads as a line.
            blocks.extend(_text_blocks(shape, in_body=False))
            continue
        table = getattr(shape, "table", None) if getattr(shape, "has_table", False) else None
        if table is not None:
            rows = tuple(tuple(clean_text(cell.text) for cell in row.cells) for row in table.rows)
            blocks.append(DeckBlock(kind="table", rows=rows))
            continue
        in_body = shape.is_placeholder and placeholder not in _SUBTITLES
        blocks.extend(_text_blocks(shape, in_body=in_body))

    notes = _slide_notes(slide)
    if notes:
        blocks.append(DeckBlock(kind="notes", text=notes))

    if title is None and blocks and blocks[0].kind in {"paragraph"} and len(blocks[0].text) <= 80:
        # No title placeholder, but the slide opens with a short plain line: that is its heading.
        title = blocks[0].text
        blocks = blocks[1:]

    layout = clean_text(slide.slide_layout.name) if slide.slide_layout is not None else None
    return normalise_slide(
        DeckSlide(
            index=index, title=title, blocks=tuple(blocks), layout=layout, deck_title=deck_title
        )
    )


def extract_deck(content: bytes, *, max_slides: int | None = None) -> ParsedDeck:
    """Read ``content`` as a ``.pptx`` into slide structures, in slide order.

    With ``max_slides``, a deck over the limit raises :class:`PptxTooManySlidesError` after its
    slides are counted and before any text is read. Raises :class:`PptxParseError` for anything
    that cannot be read - a corrupted or encrypted file, a legacy ``.ppt``, a zip bomb: the
    library's failure modes are many and none of them is the caller's to handle.
    """
    try:
        _check_archive(content)
        presentation = Presentation(io.BytesIO(content))
        slide_count = len(presentation.slides)
        if max_slides is not None and slide_count > max_slides:
            raise PptxTooManySlidesError(slide_count, max_slides)
        deck_title = clean_text(presentation.core_properties.title or "") or None
        slides: list[DeckSlide] = []
        for position, slide in enumerate(presentation.slides, start=1):
            if slide._element.get("show") == "0":
                continue
            slides.append(_read_slide(slide, index=position, deck_title=deck_title))
        return ParsedDeck(title=deck_title, slides=tuple(slides), slide_count=slide_count)
    except PptxParseError:
        raise
    except Exception as exc:  # noqa: BLE001 - see the docstring
        raise PptxParseError(str(exc)) from exc
