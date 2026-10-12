"""Real, parseable ``.pptx`` decks for tests, built in memory with python-pptx."""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

TITLE_SLIDE, TITLE_AND_CONTENT, SECTION_HEADER, TWO_CONTENT, TITLE_ONLY, BLANK = 0, 1, 2, 3, 5, 6


@dataclass
class SlideSpec:
    title: str | None = None
    layout: int = TITLE_AND_CONTENT
    subtitle: str | None = None
    bullets: list[tuple[str, int]] = field(default_factory=list)
    numbered: list[str] = field(default_factory=list)
    paragraph: str | None = None
    table: list[list[str]] | None = None
    notes: str | None = None
    hidden: bool = False


def _set_numbered(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    etree.SubElement(properties, f"{{{_A}}}buAutoNum", type="arabicPeriod")


def make_deck(slides: list[SlideSpec], *, title: str | None = None) -> bytes:
    presentation = Presentation()
    if title:
        presentation.core_properties.title = title
    for spec in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[spec.layout])
        if spec.hidden:
            slide._element.set("show", "0")
        if slide.shapes.title is not None:
            slide.shapes.title.text = spec.title or ""
        placeholders = [p for p in slide.placeholders if p.placeholder_format.idx != 0]

        if spec.layout == TITLE_SLIDE and placeholders:
            placeholders[0].text = spec.subtitle or ""
        elif spec.bullets or spec.numbered:
            body = placeholders[0]
            frame = body.text_frame
            first = True
            for text, level in spec.bullets:
                paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
                paragraph.text = text
                paragraph.level = level
                first = False
            for text in spec.numbered:
                paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
                paragraph.text = text
                _set_numbered(paragraph)
                first = False
        elif spec.layout in {TITLE_AND_CONTENT, TWO_CONTENT} and placeholders:
            # An unused body placeholder would be an empty box in the deck; remove it.
            for placeholder in placeholders:
                placeholder._element.getparent().remove(placeholder._element)

        if spec.paragraph:
            box = slide.shapes.add_textbox(Inches(1), Inches(5.2), Inches(8), Inches(1))
            box.text_frame.text = spec.paragraph
        if spec.table:
            rows, columns = len(spec.table), max(len(row) for row in spec.table)
            shape = slide.shapes.add_table(
                rows, columns, Inches(1), Inches(3), Inches(8), Emu(400_000 * rows)
            )
            for r, row in enumerate(spec.table):
                for c, cell in enumerate(row):
                    shape.table.cell(r, c).text = cell
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def sample_training_deck() -> bytes:
    """A small, realistic induction deck: a title slide, sections, bullets, numbered steps, a
    table and speaker notes - the shapes a training deck really has."""
    return make_deck(
        [
            SlideSpec(
                "Workplace Safety Induction",
                layout=TITLE_SLIDE,
                subtitle="A short guide for new employees",
            ),
            SlideSpec(
                "Agenda",
                bullets=[
                    ("Why safety matters", 0),
                    ("Common hazards", 0),
                    ("What to do in an emergency", 0),
                ],
            ),
            SlideSpec("Why Safety Matters", layout=SECTION_HEADER),
            SlideSpec(
                "Our Shared Responsibility",
                bullets=[
                    ("Safety is everyone's responsibility, not only the safety team's.", 0),
                    ("Most workplace accidents can be prevented.", 0),
                    ("Reporting a risk early protects you and your colleagues.", 0),
                ],
                notes=(
                    "Take a moment to think about the last time you noticed something unsafe at "
                    "work. Speaking up is always better than staying silent."
                ),
            ),
            SlideSpec(
                "Common Hazards",
                bullets=[
                    ("Slips, trips and falls", 0),
                    ("Wet floors", 1),
                    ("Loose cables", 1),
                    ("Electrical faults", 0),
                    ("Manual handling injuries", 0),
                ],
            ),
            SlideSpec(
                "Keeping Walkways Clear",
                bullets=[
                    ("Keep walkways and exits free of boxes and equipment.", 0),
                    ("Clean up spills immediately or put up a warning sign.", 0),
                    ("Tuck cables away so nobody can trip over them.", 0),
                ],
            ),
            SlideSpec("In an Emergency", layout=SECTION_HEADER),
            SlideSpec(
                "If You Discover a Fire",
                numbered=[
                    "Raise the alarm by pressing the nearest call point.",
                    "Call the emergency number 112.",
                    "Leave the building by the nearest safe exit.",
                    "Gather at the assembly point and wait for the Fire Marshal.",
                ],
                notes="Never use a lift during a fire. Always use the stairs.",
            ),
            SlideSpec(
                "Choosing the Right Extinguisher",
                table=[
                    ["Type", "Use it on", "Never use it on"],
                    ["Water", "Paper and wood", "Electrical equipment"],
                    ["CO2", "Electrical equipment", "Metal fires"],
                    ["Foam", "Flammable liquids", "Electrical equipment"],
                ],
            ),
            SlideSpec(
                "Personal Protective Equipment",
                bullets=[
                    ("Wear safety shoes in the warehouse.", 0),
                    ("Use gloves when handling sharp or hot items.", 0),
                    ("Wear a helmet wherever there is a risk of falling objects.", 0),
                ],
            ),
            SlideSpec(
                "Key Takeaways",
                bullets=[
                    ("Stay alert and keep your area tidy.", 0),
                    ("Know where the exits and extinguishers are.", 0),
                    ("Report every hazard and every incident.", 0),
                ],
            ),
            SlideSpec("Thank You", layout=TITLE_ONLY, notes="Questions are welcome."),
            SlideSpec("Internal draft", hidden=True, bullets=[("Do not publish this slide.", 0)]),
        ],
        title="Workplace Safety Induction",
    )
